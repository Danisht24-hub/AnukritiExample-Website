from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, EmailStr
from openpyxl import Workbook, load_workbook
from pathlib import Path
from datetime import datetime

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

EXCEL_FILE = Path(__file__).resolve().parent / "registrations.xlsx"

print(f"Excel file location: {EXCEL_FILE}")


class Registration(BaseModel):
    name: str
    roll: str
    email: EmailStr
    mobile: str


def initialize_excel():
    if not EXCEL_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Registrations"

        sheet.append([
            "Timestamp",
            "Name",
            "Roll",
            "Email",
            "Mobile",
        ])

        workbook.save(EXCEL_FILE)
        workbook.close()


initialize_excel()


@app.post("/register")
def register_user(registration: Registration):
    try:
        workbook = load_workbook(EXCEL_FILE)

        try:
            sheet = workbook["Registrations"]

            sheet.append([
                datetime.now().isoformat(timespec="seconds"),
                registration.name,
                registration.roll,
                str(registration.email),
                registration.mobile,
            ])

            workbook.save(EXCEL_FILE)

        finally:
            workbook.close()

        return {
            "success": True,
            "message": "Registration saved successfully!",
        }

    except Exception as e:
        print(f"Excel saving error: {type(e).__name__}: {e}")

        raise HTTPException(
            status_code=500,
            detail="Unable to save registration.",
        )


@app.get("/")
def home():
    return {
        "message": "Registration API is running",
        "excel_file": str(EXCEL_FILE),
    }